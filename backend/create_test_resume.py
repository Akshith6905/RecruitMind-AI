from docx import Document

doc = Document()
doc.add_heading('John Doe - Resume', 0)
doc.add_paragraph('Email: john.doe@example.com')
doc.add_paragraph('Phone: +1 555-123-4567')
doc.add_paragraph('Skills: Python, Machine Learning, FastAPI, SQL, Docker, AWS')
doc.add_heading('Experience', level=1)
doc.add_paragraph('Machine Learning Engineer at TechCorp. Experience building scalable ML APIs using FastAPI.')
doc.save('test_resume.docx')
